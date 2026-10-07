"""Acesso aos dados do PetCuida e fábrica de backends.

O backend SQLite é a implementação funcional desta entrega alfa. Os adaptadores
remotos são opcionais e documentados separadamente; não use este modo de demo
para autenticação ou dados reais.
"""
from datetime import date, datetime, timezone
import json
import re
import unicodedata
from .db import db

NOW = lambda: datetime.now(timezone.utc).isoformat()
MAX_IMPORT_ROWS = 500


def _text(value, field, *, required=False, limit=500, default=""):
    if value is None:
        value = default
    if not isinstance(value, str):
        value = str(value)
    value = value.strip()
    if required and not value:
        raise ValueError(f"Informe {field}.")
    if len(value) > limit:
        raise ValueError(f"O campo {field} excede o limite de {limit} caracteres.")
    return value


def _positive_int(value, field):
    if isinstance(value, bool):
        raise ValueError(f"{field} inválido.")
    try:
        parsed = int(value)
    except (TypeError, ValueError):
        raise ValueError(f"{field} inválido.") from None
    if parsed <= 0:
        raise ValueError(f"{field} inválido.")
    return parsed


def normalize_professional_row(row):
    if not isinstance(row, dict):
        raise ValueError("Cada linha da equipe deve ser um objeto com nome e cargo.")

    def key(value):
        normalized = unicodedata.normalize("NFKD", str(value or ""))
        normalized = "".join(char for char in normalized if not unicodedata.combining(char))
        return re.sub(r"[^a-z0-9]+", "_", normalized.lower()).strip("_")

    values = {key(k): v for k, v in row.items()}
    if not any(str(v or "").strip() for v in values.values()):
        return None
    name = next((values.get(k) for k in ("name", "nome", "professional_name", "profissional") if values.get(k)), None)
    if not name:
        raise ValueError("Há uma linha preenchida sem nome de profissional.")
    role = next((values.get(k) for k in ("role", "cargo", "especialidade", "funcao") if values.get(k)), "Profissional")
    registration = next((values.get(k) for k in ("registration", "registro", "crmv", "conselho") if values.get(k)), "")
    return {
        "name": _text(name, "o nome do profissional", required=True, limit=160),
        "role": _text(role, "o cargo/especialidade", limit=120, default="Profissional") or "Profissional",
        "registration": _text(registration, "o registro", limit=120),
    }


class SQLiteRepository:
    def __init__(self, path):
        self.path = path

    def authenticate(self, username, password):
        if not isinstance(username, str) or not isinstance(password, str):
            return None
        with db(self.path) as c:
            row = c.execute(
                "SELECT id,username,account_type,display_name,organization_id FROM users WHERE username=? AND password=?",
                (username.strip(), password),
            ).fetchone()
            return dict(row) if row else None

    def register(self, data):
        username = _text(data.get("username"), "um usuário", required=True, limit=64)
        password = data.get("password") or ""
        if not isinstance(password, str):
            raise ValueError("Senha inválida.")
        if len(username) < 3 or len(password) < 4:
            raise ValueError("Informe um usuário com pelo menos 3 caracteres e uma senha com pelo menos 4.")
        account_type = data.get("account_type", "tutor")
        if account_type not in ("tutor", "organization"):
            raise ValueError("Tipo de conta inválido.")

        with db(self.path) as c:
            if c.execute("SELECT 1 FROM users WHERE username=?", (username,)).fetchone():
                raise ValueError("Este usuário já está cadastrado.")
            organization_id = None
            display_name = _text(data.get("display_name") or username, "o nome de apresentação", limit=160)
            if account_type == "organization":
                legal_name = _text(data.get("legal_name"), "a razão social", required=True, limit=160)
                trade_name = _text(data.get("trade_name") or legal_name, "o nome de apresentação", required=True, limit=160)
                org_type = _text(data.get("org_type", "Clínica"), "o tipo de organização", required=True, limit=80)
                document = _text(data.get("document"), "a identificação", limit=40)
                location = _text(data.get("location"), "a localização", limit=160)
                cur = c.execute(
                    "INSERT INTO organizations(legal_name,trade_name,org_type,document,location,created_at) VALUES (?,?,?,?,?,?)",
                    (legal_name, trade_name, org_type, document, location, NOW()),
                )
                organization_id = cur.lastrowid
                display_name = trade_name
            cur = c.execute(
                "INSERT INTO users(username,password,account_type,display_name,organization_id,created_at) VALUES (?,?,?,?,?,?)",
                (username, password, account_type, display_name, organization_id, NOW()),
            )
            return {
                "id": cur.lastrowid,
                "username": username,
                "account_type": account_type,
                "display_name": display_name,
                "organization_id": organization_id,
            }

    def list_pets(self, user_id):
        with db(self.path) as c:
            return [dict(r) for r in c.execute("SELECT * FROM pets WHERE user_id=? ORDER BY id", (user_id,))]

    def add_pet(self, user_id, data):
        name = _text(data.get("name"), "o nome do pet", required=True, limit=100)
        species = _text(data.get("species", "Cão"), "a espécie", required=True, limit=40)
        with db(self.path) as c:
            if not c.execute("SELECT id FROM users WHERE id=?", (user_id,)).fetchone():
                raise ValueError("Usuário não encontrado.")
            cur = c.execute(
                "INSERT INTO pets(user_id,name,species,breed,birth_date,notes) VALUES (?,?,?,?,?,?)",
                (user_id, name, species, _text(data.get("breed"), "a raça", limit=100), _text(data.get("birth_date"), "a data", limit=20), _text(data.get("notes"), "as observações", limit=2000)),
            )
            return cur.lastrowid

    def services(self):
        with db(self.path) as c:
            legacy = [dict(r) for r in c.execute(
                "SELECT id,name,provider,location,price_cents,description,'legacy' offer_type,NULL organization_id,NULL professional_id,provider organization_name,provider professional_name FROM services ORDER BY id"
            )]
            offers = [dict(r) for r in c.execute(
                "SELECT o.id,o.name,o.location,o.price_cents,o.description,'offer' offer_type,o.organization_id,o.professional_id,g.legal_name organization_name,g.trade_name organization_trade_name,p.name professional_name FROM service_offers o JOIN organizations g ON g.id=o.organization_id JOIN professionals p ON p.id=o.professional_id WHERE o.status='publicado' AND p.active=1 ORDER BY o.id"
            )]
            return legacy + offers

    def appointments(self, user_id):
        sql = "SELECT a.*,p.name pet_name,s.name service_name,s.provider,s.location,s.price_cents FROM appointments a JOIN pets p ON p.id=a.pet_id JOIN services s ON s.id=a.service_id WHERE a.user_id=? ORDER BY a.scheduled_for"
        with db(self.path) as c:
            rows = [dict(r) for r in c.execute(sql, (user_id,))]
            rows += [dict(r) for r in c.execute(
                "SELECT d.*,p.name pet_name,o.name service_name,g.legal_name organization_name,g.trade_name organization_trade_name,pro.name professional_name,o.location,o.price_cents FROM demand_requests d JOIN pets p ON p.id=d.pet_id JOIN service_offers o ON o.id=d.offer_id JOIN organizations g ON g.id=o.organization_id LEFT JOIN professionals pro ON pro.id=d.accepted_professional_id WHERE d.user_id=? ORDER BY d.scheduled_for",
                (user_id,),
            )]
            return rows

    def book(self, user_id, data):
        if data.get("pet_id") in (None, "") or data.get("scheduled_for") in (None, ""):
            raise ValueError("Informe o pet e a data.")
        pet_id = _positive_int(data["pet_id"], "Pet")
        scheduled_for = _text(data["scheduled_for"], "a data", required=True, limit=10)
        try:
            date.fromisoformat(scheduled_for)
        except ValueError:
            raise ValueError("Data inválida. Use o formato AAAA-MM-DD.") from None
        with db(self.path) as c:
            if not c.execute("SELECT id FROM pets WHERE id=? AND user_id=?", (pet_id, user_id)).fetchone():
                raise ValueError("Pet inválido para este usuário.")
            if data.get("offer_id") not in (None, "") or data.get("offer_type") == "offer":
                offer_id = _positive_int(data.get("offer_id"), "Oferta")
                if not c.execute("SELECT id FROM service_offers WHERE id=? AND status='publicado'", (offer_id,)).fetchone():
                    raise ValueError("Oferta indisponível.")
                cur = c.execute(
                    "INSERT INTO demand_requests(user_id,pet_id,offer_id,scheduled_for,created_at) VALUES (?,?,?,?,?)",
                    (user_id, pet_id, offer_id, scheduled_for, NOW()),
                )
                return cur.lastrowid
            service_id = _positive_int(data.get("service_id"), "Serviço")
            if not c.execute("SELECT id FROM services WHERE id=?", (service_id,)).fetchone():
                raise ValueError("Serviço indisponível.")
            cur = c.execute(
                "INSERT INTO appointments(user_id,pet_id,service_id,scheduled_for) VALUES (?,?,?,?)",
                (user_id, pet_id, service_id, scheduled_for),
            )
            return cur.lastrowid

    def _require_organization(self, c, org_id):
        if not org_id:
            raise ValueError("Organização não informada.")
        row = c.execute("SELECT * FROM organizations WHERE id=?", (org_id,)).fetchone()
        if not row:
            raise ValueError("Organização não encontrada.")
        return row

    def organization(self, org_id):
        with db(self.path) as c:
            return dict(self._require_organization(c, org_id))

    def professionals(self, org_id):
        with db(self.path) as c:
            self._require_organization(c, org_id)
            return [dict(r) for r in c.execute(
                "SELECT * FROM professionals WHERE organization_id=? ORDER BY active DESC,name", (org_id,)
            )]

    def add_professional(self, org_id, data):
        normalized = normalize_professional_row(data)
        if normalized is None:
            raise ValueError("Informe os dados do profissional.")
        return self.add_professionals_bulk(org_id, [normalized])[0]

    def add_professionals_bulk(self, org_id, rows):
        if not rows:
            raise ValueError("O arquivo não contém profissionais válidos.")
        if len(rows) > MAX_IMPORT_ROWS:
            raise ValueError(f"O limite é de {MAX_IMPORT_ROWS} profissionais por importação.")
        normalized = [normalize_professional_row(row) for row in rows]
        normalized = [row for row in normalized if row]
        if not normalized:
            raise ValueError("O arquivo não contém profissionais válidos.")
        with db(self.path) as c:
            self._require_organization(c, org_id)
            now = NOW()
            created_ids = []
            for row in normalized:
                cursor = c.execute(
                    "INSERT INTO professionals(organization_id,name,role,registration,created_at) VALUES (?,?,?,?,?)",
                    (org_id, row["name"], row["role"], row["registration"], now),
                )
                created_ids.append(cursor.lastrowid)
            return created_ids

    def service_offers(self, org_id):
        with db(self.path) as c:
            self._require_organization(c, org_id)
            return [dict(r) for r in c.execute(
                "SELECT o.*,g.legal_name organization_name,g.trade_name organization_trade_name,p.name professional_name,p.role professional_role FROM service_offers o JOIN organizations g ON g.id=o.organization_id JOIN professionals p ON p.id=o.professional_id WHERE o.organization_id=? ORDER BY o.id DESC",
                (org_id,),
            )]

    def add_offer(self, org_id, data):
        name = _text(data.get("name"), "o serviço", required=True, limit=160)
        description = _text(data.get("description"), "a descrição", limit=2000)
        location = _text(data.get("location"), "a localização", required=True, limit=160)
        professional_id = _positive_int(data.get("professional_id"), "Profissional")
        try:
            price_cents = int(data.get("price_cents") or 0)
        except (TypeError, ValueError):
            raise ValueError("Valor inválido.") from None
        if price_cents < 0 or price_cents > 99_999_999:
            raise ValueError("O valor deve estar entre R$ 0,00 e R$ 999.999,99.")
        with db(self.path) as c:
            self._require_organization(c, org_id)
            if not c.execute(
                "SELECT id FROM professionals WHERE id=? AND organization_id=? AND active=1", (professional_id, org_id)
            ).fetchone():
                raise ValueError("Profissional inválido para esta organização.")
            cur = c.execute(
                "INSERT INTO service_offers(organization_id,professional_id,name,description,location,price_cents,created_at) VALUES (?,?,?,?,?,?,?)",
                (org_id, professional_id, name, description, location, price_cents, NOW()),
            )
            return cur.lastrowid

    def demands(self, org_id):
        sql = "SELECT d.*,u.username requester_username,p.name pet_name,o.name service_name,o.location,o.price_cents,g.legal_name organization_name,g.trade_name organization_trade_name,pro.name professional_name FROM demand_requests d JOIN users u ON u.id=d.user_id JOIN pets p ON p.id=d.pet_id JOIN service_offers o ON o.id=d.offer_id JOIN organizations g ON g.id=o.organization_id LEFT JOIN professionals pro ON pro.id=d.accepted_professional_id WHERE o.organization_id=? ORDER BY CASE d.status WHEN 'pendente' THEN 0 ELSE 1 END,d.created_at DESC"
        with db(self.path) as c:
            self._require_organization(c, org_id)
            return [dict(r) for r in c.execute(sql, (org_id,))]

    def accept_demand(self, org_id, demand_id, professional_id):
        professional_id = _positive_int(professional_id, "Profissional")
        demand_id = _positive_int(demand_id, "Demanda")
        with db(self.path) as c:
            self._require_organization(c, org_id)
            valid = c.execute(
                "SELECT d.id FROM demand_requests d JOIN service_offers o ON o.id=d.offer_id WHERE d.id=? AND o.organization_id=? AND d.status=?",
                (demand_id, org_id, "pendente"),
            ).fetchone()
            if not valid:
                raise ValueError("Demanda não está disponível.")
            if not c.execute(
                "SELECT id FROM professionals WHERE id=? AND organization_id=? AND active=1", (professional_id, org_id)
            ).fetchone():
                raise ValueError("Profissional inválido.")
            c.execute(
                "UPDATE demand_requests SET status='aceita',accepted_professional_id=? WHERE id=?",
                (professional_id, demand_id),
            )
            return demand_id


def parse_professionals_file(file_storage):
    filename = (file_storage.filename or "").lower()
    raw = file_storage.read()
    if not raw:
        raise ValueError("O arquivo está vazio.")
    if filename.endswith(".json"):
        try:
            payload = json.loads(raw.decode("utf-8-sig"))
        except (UnicodeDecodeError, json.JSONDecodeError):
            raise ValueError("JSON inválido. Salve o arquivo em UTF-8 e tente novamente.") from None
        if isinstance(payload, list):
            rows = payload
        elif isinstance(payload, dict) and isinstance(payload.get("professionals"), list):
            rows = payload["professionals"]
        else:
            raise ValueError("O JSON deve conter uma lista ou uma propriedade 'professionals' com uma lista.")
    elif filename.endswith(".xlsx"):
        try:
            from io import BytesIO
            from openpyxl import load_workbook
            workbook = load_workbook(filename=BytesIO(raw), read_only=True, data_only=True)
            sheet = workbook.active
            if sheet.max_row and sheet.max_row > MAX_IMPORT_ROWS + 1:
                raise ValueError(f"O limite é de {MAX_IMPORT_ROWS} profissionais por importação.")
            values = list(sheet.values)
            workbook.close()
        except ValueError:
            raise
        except Exception as error:
            raise ValueError("Não foi possível ler o XLSX. Confira se o arquivo está íntegro.") from error
        rows = _sheet_rows(values)
    elif filename.endswith(".xls"):
        try:
            import xlrd
            workbook = xlrd.open_workbook(file_contents=raw)
            sheet = workbook.sheet_by_index(0)
            if sheet.nrows > MAX_IMPORT_ROWS + 1:
                raise ValueError(f"O limite é de {MAX_IMPORT_ROWS} profissionais por importação.")
            values = [sheet.row_values(i) for i in range(sheet.nrows)]
        except ValueError:
            raise
        except Exception as error:
            raise ValueError("Não foi possível ler o XLS. Confira se o arquivo está íntegro.") from error
        rows = _sheet_rows(values)
    else:
        raise ValueError("Formato inválido. Envie JSON, XLS ou XLSX.")
    if not isinstance(rows, list):
        raise ValueError("A lista de profissionais está malformada.")
    if len(rows) > MAX_IMPORT_ROWS:
        raise ValueError(f"O limite é de {MAX_IMPORT_ROWS} profissionais por importação.")
    return rows


def _sheet_rows(values):
    if not values:
        return []
    headers = [str(value or "").strip() for value in values[0]]
    return [
        {headers[index]: row[index] for index in range(min(len(headers), len(row))) if headers[index]}
        for row in values[1:]
    ]


def build_repository(config):
    backend = (config.get("BACKEND") or "sqlite").lower()
    if backend == "sqlite":
        return SQLiteRepository(config["DB_PATH"])
    if backend in ("firestore", "firebase", "firebase_firestore"):
        from .data_backends.firestore_adapter import FirestoreRepository
        return FirestoreRepository.from_config(config)
    if backend == "supabase":
        from .data_backends.supabase_adapter import SupabaseRepository
        return SupabaseRepository.from_config(config)
    raise RuntimeError(f"Backend desconhecido: {backend}. Use sqlite, firestore ou supabase.")
